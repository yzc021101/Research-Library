"""Cross-platform frozen-runtime acceptance: live retrieval and upgrade-safe MCP."""
from pathlib import Path
import os,sys,json,subprocess,time,urllib.request,asyncio,shutil,zipfile
from datetime import timedelta
from xml.etree import ElementTree as ET
from mcp import ClientSession,StdioServerParameters
from mcp.client.stdio import stdio_client
from openpyxl import Workbook
from docx import Document
engine=Path(sys.argv[1]).resolve();root=Path(sys.argv[2]).resolve();out=Path(sys.argv[3]).resolve()
assert not root.exists(),'Use a fresh isolated root'
for n in ('inbox','data','logs'):(root/n).mkdir(parents=True,exist_ok=True)
inbox=root/'inbox';(root/'data/settings.json').write_text(json.dumps({'source_folders':[str(inbox)],'app_managed':True,'auto_scan':False,'search_mode':'keyword'}),encoding='utf-8')
(inbox/'anchor.txt').write_text('ExistingSearchAnchor');(inbox/'empty.txt').write_text('  \n')
w=Workbook();w.active['A1']='#DIV/0!';w.active['B1']='ErrorWorkbookAnchor';w.active['B2']=42;w.save(inbox/'error.xlsx')
env={**os.environ,'RESEARCH_LIBRARY_ROOT':str(root),'HF_HUB_OFFLINE':'1','PYTHONUTF8':'1'}
log=(root/'logs/backend.log').open('wb');backend=subprocess.Popen([str(engine),'backend'],env=env,stdout=log,stderr=log)
checks={};details={}
def req(route,body=None):
 info=json.loads((root/'data/service.json').read_text(encoding='utf-8'));request=urllib.request.Request('http://127.0.0.1:'+str(info['port'])+route,data=None if body is None else json.dumps(body).encode(),headers={'Authorization':'Bearer '+info['token'],'Content-Type':'application/json'})
 with urllib.request.urlopen(request,timeout=60) as r:return json.load(r)
def tool(name,**arguments):return req('/tool',{'name':name,'arguments':arguments})
def wait_sync():
 for _ in range(1200):
  state=req('/status')
  if not state['busy']:return state
  time.sleep(.1)
 raise TimeoutError('indexing')
async def upgrade():
 # Start at an app-style path, then move it and install another directory there.
 # Windows copies files because test users need not hold symlink privileges.
 target=root/'replaceable-engine';shutil.copytree(engine.parent,target);print('Engine copied',flush=True)
 command=[str(target/engine.name),'server'] if os.name!='nt' else [os.environ.get('COMSPEC','cmd.exe'),'/d','/c',str(root/'run-mcp.cmd')]
 async with stdio_client(StdioServerParameters(command=command[0],args=command[1:],env={**env,'RESEARCH_LIBRARY_ENGINE':str(target/engine.name)})) as (read,write):
  async with ClientSession(read,write,read_timeout_seconds=timedelta(seconds=60)) as session:
   await session.initialize();print("MCP initialized",flush=True)
   target.rename(root/'previous-engine')
   target.mkdir();shutil.copy2(engine,target/engine.name)
   # Deliberately absent modules at the old path expose accidental lazy reads.
   r=await session.call_tool('library_status',{});checks['old_connection_status_after_replace']=not r.isError
   r=await session.call_tool('search_passages',{'query':'ExistingSearchAnchor','mode':'keyword'})
   checks['old_connection_search_after_replace']=not r.isError and any('ExistingSearchAnchor' in getattr(c,'text','') for c in r.content)
   # Replacement files can be overwritten even while the bridge is connected.
   shutil.copy2(engine,target/engine.name);checks['replacement_path_writable']=True
 command=[str(engine),'server'] if os.name!='nt' else [os.environ.get('COMSPEC','cmd.exe'),'/d','/c',str(root/'run-mcp.cmd')]
 async with stdio_client(StdioServerParameters(command=command[0],args=command[1:],env={**env,'RESEARCH_LIBRARY_ENGINE':str(engine)})) as (read,write):
  async with ClientSession(read,write,read_timeout_seconds=timedelta(seconds=60)) as session:
   await session.initialize();print("MCP initialized",flush=True);r=await session.call_tool('search_documents',{'query':'anchor'});checks['new_connection_search']=not r.isError and any('anchor.txt' in getattr(c,'text','') for c in r.content)
try:
 for _ in range(200):
  try:req('/status');break
  except Exception:
   if backend.poll() is not None:raise RuntimeError('backend exited')
   time.sleep(.2)
 first=tool('sync_library');checks['empty_skipped']=any(s['state']=='skipped' and s['path'].endswith('empty.txt') for s in first['status']['sources'])
 checks['formula_error_not_failed']=any(s['state']=='ready' and s['path'].endswith('error.xlsx') for s in first['status']['sources'])
 repeat=tool('sync_library');checks['unchanged_not_reparsed']=repeat['metrics']['parsed']==0
 for i in range(12):
  d=Document()
  for n in range(150):d.add_paragraph(f'BatchDocument{i} Paragraph{n} '+'Long local research text. '*30)
  d.save(inbox/f'batch-{i}.docx')
 req('/sync',{});latencies=[];observed=[]
 for _ in range(1200):
  s=req('/status')
  if not s['busy']:break
  t=time.monotonic();r=tool('search_passages',query='ExistingSearchAnchor',mode='keyword');latencies.append(time.monotonic()-t)
  observed.append(bool(r.get('indexing') or r.get('status')=='indexing'))
  assert r.get('count',0)>0 or r.get('retryable'),r
  time.sleep(.05)
 else:raise TimeoutError('batch')
 # Indexing may finish between /status and the tool call; completed answers are valid too.
 checks['search_during_indexing']=bool(latencies) and any(observed)
 checks['indexing_responsive']=bool(latencies) and max(latencies)<3
 checks['bounded_parallelism']=1<=s['sync_progress']['metrics']['workers']<=4
 checks['batch_complete']=s['active_documents']==14 and not s.get('last_error')
 details.update(search_max_seconds=max(latencies,default=0),search_observations=len(latencies),metrics=s['sync_progress']['metrics'])
 asyncio.run(upgrade())
finally:
 try:req('/shutdown',{})
 except Exception:pass
 try:backend.wait(15)
 except subprocess.TimeoutExpired:backend.terminate();backend.wait(15)
 log.close()
out.write_text(json.dumps({'checks':checks,'details':details},ensure_ascii=False,indent=2),encoding='utf-8');print(out.read_text(encoding='utf-8'));assert checks and all(checks.values()),checks
