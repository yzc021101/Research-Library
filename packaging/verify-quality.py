"""Black-box quality regressions against an installed/frozen runtime, on both OS."""
from pathlib import Path
import os,sys,json,subprocess,time,urllib.request,zipfile,struct,hashlib
from lxml import etree as ET
from openpyxl import Workbook
from docx import Document
from pptx import Presentation
from pptx.util import Inches
engine=Path(sys.argv[1]).resolve();root=Path(sys.argv[2]).resolve();output=Path(sys.argv[3]).resolve()
assert not root.exists(), 'Use a new disposable test root'
for part in ['data','inbox','logs']:(root/part).mkdir(parents=True,exist_ok=True)
inbox=root/'inbox';checks={}
(root/'data/settings.json').write_text(json.dumps({'source_folders':[str(inbox)],'app_managed':True,'auto_scan':False,'search_mode':'keyword','keep_snapshots':False}),encoding='utf-8')
# A formula on the first row exercises both context and cell presentation.
w=Workbook();s=w.active;s['A1']='=1/4';s['A1'].number_format='0.0%';s['B1']='ModelProbe';s['A2']='=1+2';s['B2']='=1+2';s['C2']='=1+2';w.save(inbox/'model.xlsx')
with zipfile.ZipFile(inbox/'model.xlsx') as z:entries={i.filename:z.read(i) for i in z.infolist()}
ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}';sheet=ET.fromstring(entries['xl/worksheets/sheet1.xml'])
for c in sheet.iter(ns+'c'):
 v=c.find(ns+'v')
 if c.get('r')=='A1':v.text='0.25'
 if c.get('r')=='B2':c.set('t','str');v.text=''
 if c.get('r')=='C2':c.set('t','e');v.text='#REF!'
entries['xl/worksheets/sheet1.xml']=ET.tostring(sheet)
with zipfile.ZipFile(inbox/'model.xlsx','w') as z:
 for name,data in entries.items():z.writestr(name,data)
# Text boxes and alternate shapes are omitted by the stock high-level parsers.
W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}';MC='{http://schemas.openxmlformats.org/markup-compatibility/2006}'
d=Document();d.add_paragraph('WordMainProbe');p=d.add_paragraph();r=ET.SubElement(p._p,W+'r');box=ET.SubElement(r,W+'txbxContent');para=ET.SubElement(box,W+'p');run=ET.SubElement(para,W+'r');ET.SubElement(run,W+'t').text='TextboxProbe 12345';d.save(inbox/'box.docx')
p=Presentation();slide=p.slides.add_slide(p.slide_layouts[6]);shape=slide.shapes.add_textbox(0,0,Inches(3),Inches(1));shape.text='AlternateShapeProbe 67890';node=shape.element;parent=node.getparent();index=parent.index(node);alt=ET.Element(MC+'AlternateContent');choice=ET.SubElement(alt,MC+'Choice');choice.set('Requires','a');choice.append(node);parent.insert(index,alt);p.save(inbox/'alternate.pptx')
originals={p.name:hashlib.file_digest(p.open('rb'),'sha256').hexdigest() for p in inbox.iterdir()}
env={**os.environ,'RESEARCH_LIBRARY_ROOT':str(root),'PYTHONUTF8':'1','HF_HUB_OFFLINE':'1'}
log=(root/'logs/backend.log').open('wb');process=None

def req(route,body=None):
 state=json.loads((root/'data/service.json').read_text(encoding='utf-8'))
 request=urllib.request.Request(f"http://127.0.0.1:{state['port']}"+route,data=None if body is None else json.dumps(body).encode(),headers={'Authorization':'Bearer '+state['token'],'Content-Type':'application/json'})
 with urllib.request.urlopen(request,timeout=120) as r:return json.load(r)
def tool(name,**arguments):return req('/tool',{'name':name,'arguments':arguments})
def start():
 global process
 process=subprocess.Popen([str(engine),'backend'],env=env,stdout=log,stderr=log)
 for _ in range(200):
  try:return req('/status')
  except Exception:
   if process.poll() is not None:raise RuntimeError('Engine exited; see backend.log')
   time.sleep(.2)
 raise TimeoutError('Start')
def stop():
 if process and process.poll() is None:req('/shutdown',{});process.wait(20)
def retry(path=None):
 assert req('/imports/retry',{} if path is None else {'path':str(path)})['accepted']
 for _ in range(300):
  state=req('/status')
  if not state['busy']:return state
  time.sleep(.1)
 raise TimeoutError('Retry')
try:
 start();tool('sync_library');state=req('/status')
 checks['partial_separate_searchable']=state['active_documents']==3 and next(s for s in state['sources'] if s['path'].endswith('model.xlsx'))['state']=='partial'
 hit=tool('search_passages',query='ModelProbe',mode='keyword')['results'][0]
 unit=tool('read_unit',doc_id=hit['doc_id'],unit_index=hit['unit_index'])['unit'];rows={r['index']:r['cells'] for t in unit['tables'] for r in t['rows']}
 checks['saved_percent_and_value']=rows[1]['A']['text']=='25.0%' and rows[1]['A']['saved_value']==.25
 checks['no_formula_in_search_text']='=1/4' not in hit['text'] and '=1+2' not in hit['text']
 checks['blank_vs_missing']=rows[2]['B']['text']=='' and rows[2]['B']['cached_result'] and not rows[2]['A']['cached_result']
 checks['excel_error_kept']=rows[2]['C']['text']=='#REF!'
 for marker,key in [('TextboxProbe','textbox'),('AlternateShapeProbe','alternate_shape')]:checks[key]=tool('search_passages',query=marker,mode='keyword')['count']==1
 state=retry(inbox/'model.xlsx');checks['single_partial_retry']=not state.get('last_error') and any(s['state']=='partial' for s in state['sources'])
 state=retry();checks['batch_partial_retry']=not state.get('last_error') and state['active_documents']==3
 stop();state=start();checks['restart_partial_retained']=any(s['state']=='partial' for s in state['sources'])
 checks['originals_preserved']=all(hashlib.file_digest(p.open('rb'),'sha256').hexdigest()==originals[p.name] for p in inbox.iterdir())
finally:stop();log.close()
output.write_text(json.dumps({'checks':checks,'passed':sum(checks.values()),'total':len(checks)},ensure_ascii=False,indent=2),encoding='utf-8');print(output.read_text(encoding='utf-8'));assert all(checks.values()),checks
