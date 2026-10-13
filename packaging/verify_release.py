"""Black-box validation of a relocated frozen engine; no project Python is used by the product."""
from pathlib import Path
import hashlib,json,os,shutil,subprocess,sys,time,urllib.request
engine=Path(sys.argv[1]).resolve();root=Path(sys.argv[2]).resolve();evidence=Path(sys.argv[3]).resolve()
root.mkdir(parents=True,exist_ok=True)
# This root is explicitly disposable and must never be a live library.
assert 'packag' in str(root).lower() and not (root/'data/library.sqlite3').exists()
(root/'inbox').mkdir();(root/'data').mkdir();(root/'logs').mkdir()
(root/'data/settings.json').write_text(json.dumps({'source_folders':[str(root/'inbox')],'app_managed':True,'auto_scan':False}),encoding='utf-8')
env={**os.environ,'RESEARCH_LIBRARY_ROOT':str(root),'PYTHONUTF8':'1','PATH':os.environ.get('SystemRoot','C:/Windows')+'/System32' if os.name=='nt' else '/usr/bin:/bin:/usr/sbin:/sbin'}
# Probe all modern parsers with plain data, no OCR component or model present.
from docx import Document
from pptx import Presentation
from openpyxl import Workbook
import pymupdf
marker='ReleaseProbe 800 Gbps'
(root/'inbox/test.txt').write_text(marker,encoding='utf-8');(root/'inbox/test.md').write_text('# '+marker,encoding='utf-8')
(root/'inbox/test.csv').write_text('parameter,value\nReleaseProbe,800 Gbps',encoding='utf-8')
(root/'inbox/test.tsv').write_text('parameter\tvalue\nReleaseProbe\t800 Gbps',encoding='utf-8')
(root/'inbox/test.rtf').write_text(r'{\rtf1\ansi ReleaseProbe 800 Gbps}',encoding='utf-8')
d=Document();d.add_paragraph(marker);d.save(root/'inbox/test.docx')
p=Presentation();p.slides.add_slide(p.slide_layouts[6]).shapes.add_textbox(0,0,1000000,1000000).text=marker;p.save(root/'inbox/test.pptx')
w=Workbook();w.active.append(['ReleaseProbe','800 Gbps']);w.save(root/'inbox/test.xlsx')
pdf=pymupdf.open();page=pdf.new_page();page.insert_text((30,30),marker);pdf.save(root/'inbox/test.pdf');pdf.close()
originals={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in (root/'inbox').iterdir()}
checks={};process=None
log=(root/'logs/test-backend.log').open('wb')
def start():
 global process
 process=subprocess.Popen([str(engine),'backend'],env=env,stdout=log,stderr=log,cwd=root)
 for _ in range(150):
  try:
   status=req('/status');return status
  except Exception:
   if process.poll() is not None:raise RuntimeError('backend exited: '+(root/'logs/test-backend.log').read_text(errors='replace')[-3000:])
   time.sleep(.2)
 raise TimeoutError('start')
def req(route,payload=None):
 info=json.loads((root/'data/service.json').read_text(encoding='utf-8'))
 r=urllib.request.Request('http://127.0.0.1:'+str(info['port'])+route,data=json.dumps(payload).encode() if payload is not None else None,headers={'Authorization':'Bearer '+info['token'],'Content-Type':'application/json'})
 with urllib.request.urlopen(r,timeout=180) as response:return json.load(response)
def stop():
 if process and process.poll() is None:
  req('/shutdown',{});process.wait(timeout=20)
try:
 s=start();checks['empty_no_model']=s['active_documents']==0 and all(not x['downloaded'] for x in s['models'])
 result=req('/tool',{'name':'sync_library','arguments':{}});s=req('/status');checks['nine_formats']=s['active_documents']==9
 hits=req('/tool',{'name':'search_passages','arguments':{'query':'ReleaseProbe','mode':'hybrid','top_k':20}})
 checks['keyword_without_model']=hits['mode']=='keyword' and len({x['file_format'] for x in hits['results']})==9
 for hit in hits['results']:
  unit=req('/tool',{'name':'read_unit','arguments':{'doc_id':hit['doc_id'],'unit_index':hit['unit_index']}})
  assert '800 Gbps' in json.dumps(unit)
 checks['read_units']=True
 probe=subprocess.run([str(engine),'check_mcp'],env=env,capture_output=True,text=True,timeout=50)
 checks['mcp_handshake']=probe.returncode==0 and json.loads(probe.stdout)['tool_count']==7
 # Editing and deleting a source updates keyword results without a model.
 (root/'inbox/test.txt').write_text('ChangedProbe 1600 Gbps',encoding='utf-8');(root/'inbox/test.md').unlink()
 req('/tool',{'name':'sync_library','arguments':{}})
 checks['edit_delete']=req('/status')['active_documents']==8 and req('/tool',{'name':'search_passages','arguments':{'query':'ChangedProbe','mode':'keyword'}})['count']==1
 checks['no_implicit_download']=not any((root/'cache').rglob('*.onnx'))
 checks['no_venv_dependency']=not (root/'.venv').exists()
 checks['originals_preserved']=all(hashlib.sha256(p.read_bytes()).hexdigest()==originals[p.name] for p in (root/'inbox').iterdir() if p.name!='test.txt')
 finally_note=''
finally:
 stop();log.close()
evidence.write_text(json.dumps({'checks':checks,'passed':sum(checks.values()),'total':len(checks),'engine':str(engine)},ensure_ascii=False,indent=2),encoding='utf-8')
print(evidence.read_text(encoding='utf-8'))
assert all(checks.values()),checks
