"""Local OCR for OOXML embedded raster images; never follows external relationships."""
from paths import ROOT, inside
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
import hashlib
import io
import json
import posixpath
import re
import shutil
import subprocess
import threading
import time
import warnings
import zipfile

from defusedxml import ElementTree as ET
from PIL import Image, ImageOps

CACHE = ROOT / 'cache/office-ocr'
VERSION = 5
MAX_IMAGES = 500
MAX_PIXELS = 32_000_000
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
      'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      's': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
      'xdr': 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing'}

def tag(name):
    prefix, local = name.split(':')
    return '{' + NS[prefix] + '}' + local

def cancelled(event):
    if event is not None and event.is_set():
        raise InterruptedError('同步已取消；已完成的文档保留，下次同步可继续。')

@dataclass
class Picture:
    blob: bytes
    member: str
    location: dict
    crop: tuple = (0, 0, 0, 0)
    rotation: float = 0
    flip_h: bool = False
    flip_v: bool = False

def xml(z, part):
    return ET.fromstring(z.read(part))

def relationships(z, part):
    name = posixpath.join(posixpath.dirname(part), '_rels', posixpath.basename(part) + '.rels')
    if name not in z.namelist(): return {}
    result = {}
    for r in xml(z, name):
        if r.get('TargetMode', '').lower() == 'external': continue
        target = r.get('Target', '')
        resolved = posixpath.normpath(target.lstrip('/') if target.startswith('/') else posixpath.join(posixpath.dirname(part), target))
        if resolved.startswith('../') or '://' in resolved or resolved not in z.namelist(): continue
        result[r.get('Id')] = resolved
    return result

def pictures(z, part, element, location):
    rels = relationships(z, part)
    parents = {c: p for p in element.iter() for c in p}
    for node in element.iter():
        if node.tag == tag('a:blip'):
            rid = node.get(tag('r:embed'))
        elif node.tag == '{urn:schemas-microsoft-com:vml}imagedata':
            rid = node.get(tag('r:id'))
        else: continue
        member = rels.get(rid)
        if not member: continue
        parent = parents.get(node)
        crop = (0, 0, 0, 0)
        if parent is not None:
            rect = parent.find('a:srcRect', NS)
            if rect is not None: crop = tuple(int(rect.get(k, '0')) for k in ('l', 't', 'r', 'b'))
        # Find the enclosing picture's transform, without applying a group's transform twice.
        holder = parent
        while holder is not None and holder.tag.rsplit('}', 1)[-1] not in {'pic', 'sp', 'graphicFrame'}:
            holder = parents.get(holder)
        xf = holder.find('.//a:xfrm', NS) if holder is not None else None
        rot = float(xf.get('rot', '0')) / 60000 if xf is not None else 0
        yield Picture(z.read(member), member, dict(location), crop, rot,
                      xf is not None and xf.get('flipH') in {'1', 'true'},
                      xf is not None and xf.get('flipV') in {'1', 'true'})

def embedded_images(path, suffix=None):
    suffix = suffix or path.suffix.lower()
    with zipfile.ZipFile(path) as z:
        if suffix == '.docx':
            part = 'word/document.xml'
            body = xml(z, part).find('w:body', NS)
            blocks = [e for e in body if e.tag in {tag('w:p'), tag('w:tbl')}]
            for n, block in enumerate(blocks, 1):
                yield from pictures(z, part, block, {'parent_kind': 'paragraphs', 'block_start': n,
                    'block_end': n, 'label': f'正文块 {n}'})
        elif suffix == '.pptx':
            part = 'ppt/presentation.xml'; rels = relationships(z, part)
            for n, slide in enumerate(xml(z, part).findall('p:sldIdLst/p:sldId', NS), 1):
                slide_part = rels.get(slide.get(tag('r:id')))
                if slide_part:
                    yield from pictures(z, slide_part, xml(z, slide_part),
                        {'parent_kind': 'slide', 'slide': n, 'label': f'幻灯片 {n}'})
        elif suffix == '.xlsx':
            from openpyxl.utils import get_column_letter
            part = 'xl/workbook.xml'; rels = relationships(z, part)
            for sheet in xml(z, part).findall('s:sheets/s:sheet', NS):
                sheet_part = rels.get(sheet.get(tag('r:id')))
                if not sheet_part: continue
                sheet_rels = relationships(z, sheet_part)
                for drawing in xml(z, sheet_part).findall('s:drawing', NS):
                    drawing_part = sheet_rels.get(drawing.get(tag('r:id')))
                    if not drawing_part: continue
                    for anchor in xml(z, drawing_part):
                        origin = anchor.find('xdr:from', NS)
                        cell = None
                        if origin is not None:
                            row = int(origin.findtext('xdr:row', '0', NS)) + 1
                            col = int(origin.findtext('xdr:col', '0', NS)) + 1
                            cell = f'{get_column_letter(col)}{row}'
                        loc = {'parent_kind': 'cells', 'sheet': sheet.get('name'), 'anchor_cell': cell,
                               'label': sheet.get('name') + (f' · 图片锚点 {cell}' if cell else ' · 浮动图片'),
                               'hidden': sheet.get('state', 'visible') != 'visible'}
                        yield from pictures(z, drawing_part, anchor, loc)

def run_tesseract(command, event, timeout=90):
    # communicate(timeout) drains pipes and permits prompt cancellation without zombie children.
    child = subprocess.Popen(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, encoding="utf-8", errors="replace", creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    start = time.monotonic()
    try:
        while True:
            cancelled(event)
            if time.monotonic() - start > timeout: raise TimeoutError('图片 OCR 超时，可在下次同步时重试。')
            try:
                out, err = child.communicate(timeout=0.2)
                if child.returncode: raise RuntimeError('Tesseract 识别失败：' + err.strip()[-400:])
                return out
            except subprocess.TimeoutExpired: pass
    finally:
        if child.poll() is None:
            child.terminate()
            try: child.communicate(timeout=2)
            except subprocess.TimeoutExpired: child.kill(); child.communicate()

def recognize(picture, event=None):
    from cache_lock import locked
    policy = json.dumps([VERSION, picture.crop, picture.rotation, picture.flip_h, picture.flip_v, 'eng+chi_sim'])
    key = hashlib.sha256(picture.blob + policy.encode()).hexdigest()
    with locked('ocr-'+key, event):
        return _recognize(picture, event)

def _recognize(picture, event=None):
    cancelled(event)
    from office_vector import inspect_emf
    emf = inspect_emf(picture.blob)
    if emf and emf['empty']:
        return {'text': '', 'cache_hit': False, 'skip_reason': 'empty_vector_placeholder'}
    policy = json.dumps([VERSION, picture.crop, picture.rotation, picture.flip_h, picture.flip_v, 'eng+chi_sim'])
    key = hashlib.sha256(picture.blob + policy.encode()).hexdigest()
    folder = inside(CACHE / key, ROOT)
    record, preview = folder / 'ocr.json', folder / 'image.png'
    if record.exists() and preview.exists():
        try:
            saved = json.loads(record.read_text(encoding="utf-8"))
            return {**saved, 'image_path': str(preview.relative_to(ROOT)), 'cache_hit': True}
        except (OSError, ValueError): pass
    executable = shutil.which('tesseract')
    if not executable: raise RuntimeError('未找到 Tesseract；安装本地 OCR 引擎及 eng、chi_sim 语言包后重新同步。')
    vector = emf is not None
    if not vector:
        with Image.open(io.BytesIO(picture.blob)) as header: vector = header.format == 'WMF'
    if vector:
        from office_vector import rasterize
        image, source_dimensions = rasterize(picture.blob, event)
    else:
        with warnings.catch_warnings():
            warnings.simplefilter('ignore', Image.DecompressionBombWarning)
            with Image.open(io.BytesIO(picture.blob)) as original:
                source_dimensions = [original.width, original.height]
                if original.width * original.height > 128_000_000: raise ValueError('图片超过安全解码上限（1.28 亿像素）；请拆分图片后重试。')
                if original.width * original.height > MAX_PIXELS:
                    scale = (MAX_PIXELS / (original.width * original.height)) ** .5
                    original.thumbnail((max(1, int(original.width*scale)), max(1, int(original.height*scale))), Image.Resampling.LANCZOS)
                image = ImageOps.exif_transpose(original).convert('RGBA')
    w, h = image.size
    l, t, r, b = picture.crop
    box = (round(w*l/100000), round(h*t/100000), round(w*(1-r/100000)), round(h*(1-b/100000)))
    if box[2] <= box[0] or box[3] <= box[1]: raise ValueError('图片裁剪区域无效。')
    # OOXML negative crop values add transparent margins (LibreOffice also emits
    # tiny negative values when rounding legacy Office image dimensions).
    if (box[2]-box[0]) * (box[3]-box[1]) > MAX_PIXELS:
        raise ValueError('图片外扩区域超过 3200 万像素，已跳过 OCR。')
    image = image.crop(box)
    if picture.flip_h: image = ImageOps.mirror(image)
    if picture.flip_v: image = ImageOps.flip(image)
    if picture.rotation: image = image.rotate(-picture.rotation, expand=True, resample=Image.Resampling.BICUBIC)
    white = Image.new('RGBA', image.size, 'white'); white.alpha_composite(image); image = white.convert('RGB')
    folder.mkdir(parents=True, exist_ok=True)
    displayed = image.copy(); displayed.thumbnail((2048, 2048))
    displayed.save(preview)
    work = ImageOps.autocontrast(ImageOps.grayscale(image))
    scale = min(2.0, 3000 / max(work.size))
    if scale != 1: work = work.resize((max(1, round(work.width*scale)), max(1, round(work.height*scale))), Image.Resampling.LANCZOS)
    work = ImageOps.expand(work, border=20, fill='white')
    temp = folder / f'input-{threading.get_ident()}.png'
    work.save(temp, dpi=(300, 300))
    try:
        text = run_tesseract([executable, str(temp), 'stdout', '-l', 'eng+chi_sim', '--psm', '3'], event).strip()
        if not text:
            text = run_tesseract([executable, str(temp), 'stdout', '-l', 'eng+chi_sim', '--psm', '6'], event).strip()
        text = re.sub(r'(?<=[\u3400-\u9fff])[ \t]+(?=[\u3400-\u9fff])', '', text)
        cancelled(event)
        saved = {'text': text, 'engine': 'tesseract', 'languages': 'eng+chi_sim', 'version': VERSION,
                 'source_sha256': hashlib.sha256(picture.blob).hexdigest(), 'source_dimensions': source_dimensions,
                 'preview_note': '嵌入图片预览（应用图片裁剪与旋转，最长边 2048 像素）；不代表整页排版。'}
        tmp = record.with_suffix(f'.{threading.get_ident()}.tmp')
        tmp.write_text(json.dumps(saved, ensure_ascii=False), encoding="utf-8"); tmp.replace(record)
        return {**saved, 'image_path': str(preview.relative_to(ROOT)), 'cache_hit': False}
    finally: temp.unlink(missing_ok=True)

def extract_images(path, progress=None, event=None, suffix=None):
    # Import here to avoid a circular module dependency.
    from document_formats import Unit
    units, errors = [], []
    summary = {'images': 0, 'recognized': 0, 'empty': 0, 'failed': 0, 'cache_hits': 0, 'empty_placeholders': 0}
    for n, picture in enumerate(embedded_images(path, suffix), 1):
        cancelled(event)
        if n > MAX_IMAGES:
            errors.append('内嵌图片超过 500 张，只识别前 500 张。'); summary['failed'] += 1; break
        summary['images'] += 1
        label = picture.location['label'] + f' · 图片 {n}'
        if progress: progress(stage='ocr', image=n, location=label)
        try:
            result = recognize(picture, event)
            if result.get('skip_reason') == 'empty_vector_placeholder': summary['empty_placeholders'] += 1
            summary['cache_hits'] += int(result['cache_hit'])
            if not result['text']: summary['empty'] += 1; continue
            summary['recognized'] += 1
            location = {**picture.location, 'kind': 'embedded_image', 'image_index': n, 'label': label,
                        'package_member': picture.member, 'image_path': result['image_path'],
                        'source_sha256': result['source_sha256'], 'preview_note': result['preview_note']}
            units.append(Unit('[OCR 识别文字 · 请核对图片中的数字、单位和表格关系]\n' + result['text'], location))
        except InterruptedError: raise
        except Exception as exc:
            summary['failed'] += 1
            errors.append(f'{label}：{exc}')
    return units, errors, summary
