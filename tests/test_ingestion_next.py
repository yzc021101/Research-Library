import os, sys, tempfile, unittest, threading, time, zipfile
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree as ET
fixture=tempfile.TemporaryDirectory(); os.environ['RESEARCH_LIBRARY_ROOT']=fixture.name
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'src'))
from paths import INBOX
from library import Library
from document_formats import extract_document
from openpyxl import Workbook
from pptx import Presentation
import pymupdf

class NextImport(unittest.TestCase):
 def setUp(self):
  for f in INBOX.iterdir(): f.unlink()
  self.lib=Library(lambda:{'source_folders':[str(INBOX)],'keep_snapshots':False,'search_mode':'keyword','top_k':8})
  for table in ['sources','documents','pages','chunks','chunks_fts','source_fingerprints','embedding_cache']:
   self.lib.db.execute('DELETE FROM '+table)
  self.lib.db.commit()
 def tearDown(self): self.lib.db.close()
 def test_empty_and_error_values(self):
  (INBOX/'empty.txt').write_text('  \n')
  p=INBOX/'errors.xlsx';w=Workbook();w.active['A1']='=1/0';w.active['B1']='RetainedValue';w.active['B2']=42;w.save(p)
  with zipfile.ZipFile(p) as z: parts={n:z.read(n) for n in z.namelist()}
  ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}';xml=ET.fromstring(parts['xl/worksheets/sheet1.xml'])
  c=next(c for c in xml.iter(ns+'c') if c.get('r')=='A1');c.set('t','e');c.find(ns+'v').text='#DIV/0!';parts['xl/worksheets/sheet1.xml']=ET.tostring(xml)
  with zipfile.ZipFile(p,'w') as z:
   for n,b in parts.items():z.writestr(n,b)
  result=self.lib.sync(); states={Path(r['path']).name:r['state'] for r in result['status']['sources']}
  self.assertEqual(states,{'empty.txt':'skipped','errors.xlsx':'ready'},result)
  self.assertEqual(self.lib.search('RetainedValue',mode='keyword')['count'],1)
  self.assertEqual(self.lib.sync(retry=True)['results'],[])
 def test_incremental_and_live_search(self):
  p=INBOX/'first.txt';p.write_text('ExistingAnchor')
  self.assertEqual(self.lib.sync()['metrics']['parsed'],1)
  self.assertEqual(self.lib.sync()['metrics']['parsed'],0)
  p.write_text('ChangedAnchor')
  result=self.lib.sync();self.assertEqual(result['metrics']['parsed'],1)
  self.assertEqual(self.lib.search('ExistingAnchor',mode='keyword')['count'],0)
  started=threading.Event();release=threading.Event()
  import ingest_jobs
  real=ingest_jobs.extract
  def delayed(*args):started.set();release.wait(10);return real(*args)
  (INBOX/'second.txt').write_text('SecondAnchor')
  with patch.object(ingest_jobs,'extract',side_effect=delayed):
   t=threading.Thread(target=self.lib.sync);t.start();self.assertTrue(started.wait(10))
   tick=time.monotonic();self.assertEqual(self.lib.search('ChangedAnchor',mode='keyword')['count'],1);self.assertLess(time.monotonic()-tick,2)
   release.set();t.join(20);self.assertFalse(t.is_alive())
 def test_title_and_pdf_notices(self):
  p=INBOX/'Actual Research.pptx';r=Presentation();r.core_properties.title='Stale Template MR';s=r.slides.add_slide(r.slide_layouts[6]);s.shapes.add_textbox(0,0,1000000,1000000).text='Actual Research';r.save(p)
  meta,units=extract_document(p);self.assertEqual(meta['title'],p.stem);self.assertEqual(meta['original_title'],'Stale Template MR');units.close()
  from pdf_text import search_text
  doc=pymupdf.open();page=doc.new_page();page.insert_text((50,50),'KeepNumber 42')
  # CJK font supplies an explicit usage notice without raster/OCR ambiguity.
  page.insert_text((50,100),'请勿传播与转载',fontname='china-s')
  text,notices=search_text(page,page.get_text());self.assertIn('42',text);self.assertNotIn('请勿',text);self.assertTrue(notices)
  doc.close()

if __name__=='__main__': unittest.main()
