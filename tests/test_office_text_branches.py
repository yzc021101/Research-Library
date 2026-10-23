import os,sys,tempfile,unittest,hashlib
from pathlib import Path
from lxml import etree
from docx import Document
from pptx import Presentation
from pptx.util import Inches
fixture=tempfile.TemporaryDirectory();os.environ['RESEARCH_LIBRARY_ROOT']=fixture.name
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'src'))
from document_formats import word_units,presentation_units
MC='{http://schemas.openxmlformats.org/markup-compatibility/2006}'
W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

class TextBranches(unittest.TestCase):
    def test_pptx_alternate_shape_once_and_original_unchanged(self):
        p=Path(fixture.name)/'branch.pptx';deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6])
        shape=slide.shapes.add_textbox(0,0,Inches(3),Inches(1));shape.text='PrimaryShapeAnchor'
        element=shape.element;parent=element.getparent();index=parent.index(element)
        alternate=etree.Element(MC+'AlternateContent');choice=etree.SubElement(alternate,MC+'Choice');choice.set('Requires','a')
        choice.append(element);fallback=etree.SubElement(alternate,MC+'Fallback')
        copy=etree.fromstring(etree.tostring(element))
        for t in copy.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}t'):t.text='FallbackMustNotDuplicate'
        fallback.append(copy);parent.insert(index,alternate);deck.save(p)
        original=p.read_bytes();units,_,_=presentation_units(p);text='\n'.join(u.text for u in units)
        self.assertEqual(text.count('PrimaryShapeAnchor'),1);self.assertNotIn('FallbackMustNotDuplicate',text)
        self.assertEqual(p.read_bytes(),original)
    def test_word_textbox_only_paragraph_and_table(self):
        p=Path(fixture.name)/'textbox.docx';doc=Document();doc.add_paragraph('MainTextAnchor')
        paragraph=doc.add_paragraph();run=etree.SubElement(paragraph._p,W+'r');box=etree.SubElement(run,W+'txbxContent')
        para=etree.SubElement(box,W+'p');r=etree.SubElement(para,W+'r');etree.SubElement(r,W+'t').text='FloatingTextboxAnchor'
        table=doc.add_table(1,1);table.cell(0,0).text='TableAnchor';doc.save(p)
        original=p.read_bytes();units,_,_=word_units(p);text='\n'.join(u.text for u in units)
        for anchor in ['MainTextAnchor','FloatingTextboxAnchor','TableAnchor']:self.assertIn(anchor,text)
        self.assertEqual(text.count('FloatingTextboxAnchor'),1);self.assertEqual(p.read_bytes(),original)

if __name__=='__main__':unittest.main()
