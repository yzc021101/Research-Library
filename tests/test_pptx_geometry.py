"""PPTX reading order and content retention when optional geometry is absent.

Run against either backend with DOCUMENT_FORMATS_SOURCE=src or windows/backend.
Fixtures are synthetic, and the parser must leave their bytes unchanged.
"""
import hashlib
import os
from pathlib import Path
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / os.environ.get('DOCUMENT_FORMATS_SOURCE', 'src')))

from document_formats import presentation_units
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches


class PresentationGeometryTests(unittest.TestCase):
    def setUp(self):
        output = Path(os.environ.get('PPTX_TEST_ROOT', ROOT / 'results' / 'pptx-geometry'))
        output.mkdir(parents=True, exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(dir=output)
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / 'geometry.pptx'

    def parse(self, deck):
        deck.save(self.path)
        before = hashlib.sha256(self.path.read_bytes()).digest()
        result = presentation_units(self.path)
        self.assertEqual(hashlib.sha256(self.path.read_bytes()).digest(), before)
        return result[0]

    @staticmethod
    def textbox(shapes, text, left=0, top=0):
        shape = shapes.add_textbox(left, top, Inches(2), Inches(1))
        shape.text = text
        return shape

    def test_orphan_placeholder_keeps_text_and_slide_location(self):
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = 'InheritedTitle'
        body = slide.placeholders[1]
        body.text = 'OrphanPlaceholderText'
        # No matching layout placeholder, as in exported/edited presentations.
        body._element.xpath('.//p:ph')[0].set('idx', '999')
        self.assertIsNone(body.top)
        self.assertIsNone(body.left)
        self.assertIsNone(body.width)
        self.assertIsNone(body.height)
        self.assertIsNotNone(slide.shapes.title.top)
        self.textbox(slide.shapes, 'OrdinaryText', top=Inches(4))
        units = self.parse(deck)
        self.assertEqual(units[0].location['slide'], 1)
        self.assertEqual(units[0].text.splitlines(),
                         ['InheritedTitle', 'OrdinaryText', 'OrphanPlaceholderText'])

    def test_missing_positions_and_sizes_do_not_change_valid_reading_order(self):
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        first = self.textbox(slide.shapes, 'UnpositionedFirst')
        first._element.spPr.remove(first._element.spPr.xfrm)
        self.textbox(slide.shapes, 'LowerText', top=Inches(3))
        zero = self.textbox(slide.shapes, 'ZeroPosition')
        zero._element.spPr.xfrm.remove(zero._element.spPr.xfrm.find(qn('a:ext')))
        self.assertIsNone(zero.width)
        self.assertIsNone(zero.height)
        self.textbox(slide.shapes, 'EqualFirst', left=Inches(1), top=Inches(1))
        self.textbox(slide.shapes, 'EqualSecond', left=Inches(1), top=Inches(1))
        last = self.textbox(slide.shapes, 'UnpositionedSecond')
        last._element.spPr.remove(last._element.spPr.xfrm)
        units = self.parse(deck)
        self.assertEqual(units[0].text.splitlines(), [
            'ZeroPosition', 'EqualFirst', 'EqualSecond', 'LowerText',
            'UnpositionedFirst', 'UnpositionedSecond'])
        self.assertEqual(presentation_units(self.path)[0], units)

    def test_unpositioned_group_preserves_nested_content_tables_charts_and_notes(self):
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        group = slide.shapes.add_group_shape()
        self.textbox(group.shapes, 'GroupText')
        nested = group.shapes.add_group_shape()
        child = self.textbox(nested.shapes, 'NestedText')
        child._element.spPr.remove(child._element.spPr.xfrm)
        group._element.grpSpPr.remove(group._element.grpSpPr.xfrm)
        self.assertIsNone(group.top)
        table = slide.shapes.add_table(1, 2, 0, Inches(2), Inches(4), Inches(1)).table
        table.cell(0, 0).text = 'TableAnchor'
        table.cell(0, 1).text = '88 | 99\nTwoLines'
        data = CategoryChartData()
        data.categories = ['CategoryAnchor']
        data.add_series('SeriesAnchor', (42,))
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 0, Inches(4),
                               Inches(4), Inches(2), data)
        slide.notes_slide.notes_text_frame.text = 'NotesAnchor'
        units = self.parse(deck)
        text = '\n'.join(u.text for u in units)
        for anchor in ['GroupText', 'NestedText', 'TableAnchor',
                       'SeriesAnchor', 'CategoryAnchor: 42.0', 'NotesAnchor']:
            self.assertIn(anchor, text)
            self.assertEqual(text.count(anchor), 1)
        cells = units[0].tables[0]['rows'][0]['cells']
        self.assertEqual(cells['2']['text'], '88 | 99\nTwoLines')

    def test_empty_slides_keep_original_slide_numbers(self):
        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[6])
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        missing = self.textbox(slide.shapes, 'SecondSlideText')
        missing._element.spPr.remove(missing._element.spPr.xfrm)
        units = self.parse(deck)
        self.assertEqual([u.location['slide'] for u in units], [1, 2])
        self.assertEqual([u.text for u in units], ['', 'SecondSlideText'])


if __name__ == '__main__':
    unittest.main()
