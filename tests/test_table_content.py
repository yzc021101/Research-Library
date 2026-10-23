"""Regression coverage for real cell boundaries, index/API payloads and invalid XLSX styles."""
from pathlib import Path
import os, sys, json, hashlib, zipfile, io
from datetime import datetime
LAB=Path(__file__).resolve().parents[1]
BASE=Path(os.environ.get('TABLE_TEST_ROOT',str(LAB/'results/tables-0.6.2/test-library')))
os.environ['RESEARCH_LIBRARY_ROOT']=str(BASE)
sys.path.insert(0,str(LAB/'src'))
from document_formats import extract_document, spreadsheet_units
from excel_compat import compatible_workbook
from library import Library
from openpyxl import Workbook
from openpyxl.styles import Alignment
from table_content import display_content
from docx import Document
from pptx import Presentation
from pptx.util import Inches
S=BASE/'inbox';S.mkdir(parents=True,exist_ok=True)
checks={}
def check(name,ok):
    checks[name]=bool(ok)
    assert ok,name
book=Workbook();ws=book.active;ws.title='模型更新';ws['A2']='旗舰模型更新与应用变化'
ws.append(['标题'])
values=['O12','OpenAI','GPT-6 Sol',datetime(2026,9,22),'成本、速度、自动化经济性','长文本，含 | 竖线\n及换行。'*12,'评测口径不同','不能认为能力上限突破',None,'https://example.test','未复核']
for c,v in enumerate(values,1):ws.cell(17,c,v)
ws['D17'].number_format='yyyy-mm-dd';ws['C17'].alignment=Alignment(vertical='top')
ws['A18']='SecondRowAnchor';ws['B18']='=1+2';ws['C18']=.15;ws['C18'].number_format='0.0%'
book.save(S/'table.xlsx')
raw=(S/'table.xlsx').read_bytes()
def mutate(name,entry,replacement):
    output=io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as source,zipfile.ZipFile(output,'w') as dest:
        for item in source.infolist():dest.writestr(item,replacement(source.read(item)) if item.filename==entry else source.read(item))
    (S/name).write_bytes(output.getvalue())
mutate('bad-style.xlsx','xl/styles.xml',lambda b:b.replace(b'vertical="top"',b'vertical="left"'))
mutate('bad-sheet.xlsx','xl/worksheets/sheet1.xml',lambda b:b[:-20])
mutate('bad-style-xml.xlsx','xl/styles.xml',lambda b:b[:-20])
fixed,warnings=compatible_workbook(S/'bad-style.xlsx')
with zipfile.ZipFile(io.BytesIO(fixed)) as z,zipfile.ZipFile(S/'bad-style.xlsx') as old:
    check('repair_only_styles_entry',all(z.read(n)==old.read(n) for n in old.namelist() if n!='xl/styles.xml'))
check('repair_disclosed',bool(warnings))
for name in ['bad-sheet.xlsx','bad-style-xml.xlsx']:
    try:extract_document(S/name);failed=False
    except Exception:failed=True
    check('genuinely_broken_fails_'+name,failed)
(S/'table.csv').write_text('公司,备注\n甲公司,"CsvGrid 含 | 竖线,和\n换行"\n',encoding='utf-8')
doc=Document();doc.add_paragraph('WordBeforeAnchor')
t=doc.add_table(rows=2,cols=2);t.cell(0,0).text='公司';t.cell(0,1).text='备注';t.cell(1,0).text='WordGrid';t.cell(1,1).text='保留 | 分隔符\n和换行'
doc.add_paragraph('WordAfterAnchor');doc.save(S/'table.docx')
deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6]);slide.shapes.add_textbox(Inches(0), Inches(0),Inches(3),Inches(1)).text='SlideProseAnchor'
t=slide.shapes.add_table(2,2, Inches(0),Inches(2),Inches(6),Inches(2)).table
t.cell(0,0).text='指标';t.cell(1,0).text='SlideGrid';t.cell(1,1).text='88 | 99\n多行'
slide.notes_slide.notes_text_frame.text='SlideNotesAnchor';deck.save(S/'table.pptx')
# A long row spans read units and search chunks; do not lose its actual cell structure.
wb=Workbook();wb.active['A1']='LongGrid '+('长内容 | 多行\n'*1800);wb.save(S/'long.xlsx')
original={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in S.iterdir()}
lib=Library(lambda:{'source_folders':[str(S)],'top_k':8,'search_mode':'keyword','keep_snapshots':False})
actions={Path(r['file']).name:r for r in lib.sync()['results']}
check('import_good_formats',all(actions[n]['action'] in {'indexed','unchanged_or_duplicate'} for n in ['table.xlsx','bad-style.xlsx','table.csv','table.docx','table.pptx','long.xlsx']))
for query,name in [('Sol','table.xlsx'),('CsvGrid','table.csv'),('WordGrid','table.docx'),('SlideGrid','table.pptx'),('LongGrid','long.xlsx')]:
    doc=next(d for d in lib.active_docs().values() if d['file_name']==name)
    result=lib.search(query,mode='keyword',doc_id=doc['doc_id'])['results']
    check('search_table_'+name,bool(result) and bool(result[0]['tables']))
    hit=result[0];unit=lib.read_unit(hit['doc_id'],hit['unit_index'])['unit']
    check('read_table_'+name,bool(unit['tables']))
    check('no_flattened_table_prose_'+name,'[表格行' not in hit['display_text'] and '首个非空行参考:' not in hit['display_text'])
    if name=='table.xlsx':
        row=next(r for t in hit['tables'] for r in t['rows'] if r['index']==17)
        check('dates_readable',row['cells']['D']['text']=='2026-09-22')
        check('cell_delimiters_preserved',row['cells']['F']['text']==values[5])
        check('blank_stays_blank','I' not in row['cells'])
        (BASE.parent/'sample-search.json').write_text(json.dumps(hit,ensure_ascii=False,indent=2))
    if name=='table.pptx':
        check('slide_narrative_retained','SlideProseAnchor' in unit['display_text'] and 'SlideNotesAnchor' in unit['display_text'])
# Dates, formula text and number formats must match the valid workbook despite style recovery.
a,_=extract_document(S/'table.xlsx');b,units=extract_document(S/'bad-style.xlsx')
row=next(r for u in units for t in u.tables for r in t['rows'] if r['index']==18)
check('formula_not_recalculated',row['cells']['B']['text']=='缺少已保存的计算结果' and row['cells']['B']['cached_result'] is False and row['cells']['B']['formula']=='=1+2')
check('number_format_preserved',row['cells']['C']['number_format']=='0.0%')
check('all_original_bytes_unchanged',all(hashlib.sha256((S/n).read_bytes()).hexdigest()==sha for n,sha in original.items()))
# Legacy index entries must be reindexed once and then remain stable.
doc=next(d for d in lib.active_docs().values() if d['file_name']=='table.xlsx')
m=json.loads(lib.db.execute('SELECT metadata FROM documents WHERE doc_id=?',(doc['doc_id'],)).fetchone()[0]);m['parser_version']=3
lib.db.execute('UPDATE documents SET metadata=? WHERE doc_id=?',(json.dumps(m),doc['doc_id']));lib.db.commit()
r={Path(x['file']).name:x for x in lib.sync()['results']};check('old_index_upgraded',r['table.xlsx']['action']=='indexed')
r={Path(x['file']).name:x for x in lib.sync()['results']};check('no_reindex_loop',r['table.xlsx']['action']=='unchanged_or_duplicate')
lib.db.close()
report={'passed':sum(checks.values()),'total':len(checks),'checks':checks}
(BASE.parent/'checks.json').write_text(json.dumps(report,indent=2));print(json.dumps(report,indent=2))
