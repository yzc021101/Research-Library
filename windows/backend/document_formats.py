"""Read-only format adapters. Locations describe source structure, never invented pages.

Modern formats use direct parsers. Legacy Office uses an isolated local converter.
Macros and document external-link updates are disabled for conversion.
Each unit is bounded so MCP can expand context without returning an entire workbook.
"""
from dataclasses import dataclass, field
from pathlib import Path
import csv
import io
import re
import zipfile
import time
from spooled import Spool

from table_content import cell as table_cell, table as make_table, slice_tables
from excel_compat import open_compatible_books, formula_text

from document_metadata import SCHEMA_VERSION, cover_date, institution_from_cover

SUPPORTED = {'.pdf', '.txt', '.md', '.markdown', '.docx', '.pptx', '.xlsx', '.csv', '.tsv', '.rtf'}
LEGACY = {'.doc', '.ppt', '.xls'}
SUPPORTED |= LEGACY
MAX_BYTES = 100 * 1024 * 1024
MAX_TEXT = 8_000_000
MAX_UNITS = 10000
MAX_CELLS = 2_000_000
UNIT_SIZE = 12000
PARSER_VERSION = 6

@dataclass
class Unit:
    text: str
    location: dict
    warnings: list = field(default_factory=list)
    tables: list = field(default_factory=list)

def decode_text(raw):
    if raw.startswith((b'\xff\xfe', b'\xfe\xff')):
        return raw.decode('utf-16'), 'utf-16'
    if b'\x00' in raw:
        raise ValueError('文本包含二进制数据；UTF-16 文件请保存为带 BOM 的格式。')
    for encoding in ('utf-8-sig', 'gb18030', 'big5'):
        try: return raw.decode(encoding), encoding
        except UnicodeDecodeError: pass
    raise ValueError('无法确定文本编码，请另存为 UTF-8 后重试。')

def check_package(path):
    # Reject encrypted or excessive packages before handing them to XML parsers.
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 20000 or sum(e.file_size for e in entries) > 256 * 1024 * 1024:
            raise ValueError('Office 文件解压规模超过限制（256 MB / 20000 项）。')
        if any(e.flag_bits & 1 for e in entries):
            raise ValueError('不支持加密的 Office 文档，请使用未加密副本。')
        for e in entries:
            if e.filename.endswith(('.xml', '.rels')):
                xml = archive.read(e)
                if b'<!DOCTYPE' in xml.upper() or b'<!ENTITY' in xml.upper():
                    raise ValueError('文档包含不支持的 XML 实体声明。')

def text_units(text, kind='lines'):
    # Exact physical line coordinates are retained, including blank lines.
    lines = text.splitlines()
    units, parts, size, start = [], [], 0, 1
    for n, line in enumerate(lines, 1):
        if parts and size + len(line) > UNIT_SIZE:
            units.append(Unit('\n'.join(parts), {'kind': kind, 'line_start': start, 'line_end': n-1,
                                                 'label': f'第 {start}–{n-1} 行'}))
            parts, size, start = [], 0, n
        parts.append(line); size += len(line) + 1
    if parts:
        units.append(Unit('\n'.join(parts), {'kind': kind, 'line_start': start, 'line_end': len(lines),
                                             'label': f'第 {start}–{len(lines)} 行'}))
    return units

def word_units(path):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    doc = Document(path)
    units, lines, size = [], [], 0
    paragraph, table_number, first, last, heading = 0, 0, None, None, ''
    grid_rows = []
    def flush():
        nonlocal lines, size, first, last, grid_rows
        if lines:
            label = f'正文块 {first}–{last}' + (f' · {heading}' if heading else '')
            units.append(Unit('\n'.join(lines), {'kind': 'paragraphs', 'block_start': first,
                'block_end': last, 'heading': heading or None, 'label': label},
                tables=[make_table(f'表格 {table_number}', grid_rows)] if grid_rows else []))
        lines, size, first, last, grid_rows = [], 0, None, None, []
    for block_no, block in enumerate(doc.iter_inner_content(), 1):
        if isinstance(block, Paragraph):
            paragraph += 1
            if not block.text.strip(): continue
            if block.style and block.style.name.startswith('Heading'):
                flush(); heading = block.text
            value = f'[段落 {paragraph}] {block.text}'
        elif isinstance(block, Table):
            flush(); table_number += 1
            rows = [[c.text for c in row.cells] for row in block.rows]
            if not rows or not any(value.strip() for row in rows for value in row): continue
            header = rows[0]
            for r, row in enumerate(rows, 1):
                value = f'[表格 {table_number}，行 {r}] ' + ' | '.join(
                    f'列 {c+1}（{header[c] if c < len(header) else ""}）: {v}' for c, v in enumerate(row))
                if size + len(value) > UNIT_SIZE: flush()
                first = first or block_no; last = block_no
                grid_rows.append({'index': r, 'cells': {str(c+1): table_cell(v) for c, v in enumerate(row)},
                    '_start': size, '_end': size + len(value)})
                lines.append(value); size += len(value) + 1
            flush(); continue
        else: continue
        if size + len(value) > UNIT_SIZE: flush()
        first = first or block_no; last = block_no
        lines.append(value); size += len(value) + 1
    flush()
    props = {'title': doc.core_properties.title, 'author': doc.core_properties.author}
    warnings = ['Word 按正文顺序提取；页眉页脚、修订、文本框和图片可能不完整，引用使用正文块而非排版页码。']
    return units, props, warnings

def presentation_units(path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    deck = Presentation(path)
    units = []
    def shape_content(shape):
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for child in shape.shapes: yield from shape_content(child)
        if shape.has_text_frame: yield shape.text, None
        if shape.has_table:
            for n, row in enumerate(shape.table.rows, 1):
                cells = {str(i+1): table_cell(c.text) for i, c in enumerate(row.cells)}
                text = f'[表格行 {n}] ' + ' | '.join(f'列 {i+1}: {c.text}' for i, c in enumerate(row.cells))
                yield text, {'index': n, 'cells': cells, 'table_id': shape.shape_id}
        if shape.has_chart:
            try:
                chart = shape.chart
                categories = [c.label for c in chart.plots[0].categories]
                for series in chart.series:
                    yield '[图表数据] ' + series.name + ': ' + ' | '.join(
                        f'{categories[i] if i < len(categories) else i+1}: {v}' for i, v in enumerate(series.values)), None
            except (AttributeError, ValueError, TypeError):
                yield '[此图表未能提取数据，请打开原文件核对]', None
    for i, slide in enumerate(deck.slides, 1):
        values, tables, position = [], {}, 0
        for shape in sorted(slide.shapes, key=lambda s: (s.top, s.left)):
            for text, row in shape_content(shape):
                if row:
                    table_id = row.pop('table_id')
                    row.update(_start=position, _end=position+len(text))
                    tables.setdefault(table_id, []).append(row)
                values.append(text); position += len(text)+1
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes: values.append('[演讲者备注]\n' + notes)
        text = '\n'.join(values)
        grids = [make_table(f'幻灯片 {i} · 表格 {n}', rows) for n, rows in enumerate(tables.values(), 1)]
        offset = 0
        for part in text_units(text) or [Unit('', {})]:
            units.append(Unit(part.text, {'kind': 'slide', 'slide': i, 'label': f'幻灯片 {i}',
                'slide_line_start': part.location.get('line_start', 1)},
                tables=slice_tables(grids, offset, offset+len(part.text))))
            offset += len(part.text)+1
    return units, {'title': deck.core_properties.title, 'author': deck.core_properties.author}, [
        '提取幻灯片文字、表格、可读取的图表数据和备注；图片、SmartArt 与复杂图形需要打开原文件核对。']

def reading_limits(options=None):
    large = (options or {}).get('large_file_mode', False)
    return (5_000_000, 80_000_000, 50000, 900) if large else (MAX_CELLS, 32_000_000, 20000, 300)


def spreadsheet_units(path, options=None, cancel=None, progress=None):
    from datetime import datetime, date, time as dt_time
    from openpyxl.utils import get_column_letter, column_index_from_string
    from excel_compat import sparse_rows
    from office_ocr import cancelled
    units = Spool()
    formula_book, cached_book, recovery_warnings = open_compatible_books(path)
    cell_limit, text_limit, unit_limit, seconds = reading_limits(options)
    count = text_count = 0
    deadline = time.monotonic() + seconds
    stopped = False
    try:
        props = {'title': formula_book.properties.title, 'author': formula_book.properties.creator}
        for sheet, cached in zip(formula_book.worksheets, cached_book.worksheets):
            parts, size, first, last, grid_rows = [], 0, None, None, []
            context = ''
            def flush():
                nonlocal parts, size, first, last, grid_rows
                if parts:
                    columns = sorted({key for r in grid_rows for key in r['cells']}, key=lambda k: (len(k), k))
                    last_col = columns[-1] if columns else 'A'
                    location = {'kind': 'cells', 'sheet': sheet.title, 'row_start': first, 'row_end': last,
                        'range': f'A{first}:{last_col}{last}', 'label': f'{sheet.title} · 行 {first}–{last}',
                        'hidden': sheet.sheet_state != 'visible'}
                    if columns and column_index_from_string(last_col) <= 256:
                        columns = [get_column_letter(i) for i in range(1, column_index_from_string(last_col)+1)]
                    units.append(Unit('\n'.join(parts), location, tables=[make_table(sheet.title, grid_rows,
                        [{'key': key, 'label': key} for key in columns], context)]))
                parts, size, first, last, grid_rows = [], 0, None, None, []
            for (row_no, row), (cached_no, values) in zip(sparse_rows(sheet, cancel), sparse_rows(cached, cancel), strict=True):
                cancelled(cancel)
                if row_no != cached_no: raise ValueError('公式与结果行号不一致，已停止读取。')
                reason = None
                if count + len(row) > cell_limit: reason = f'有效单元格超过 {cell_limit:,} 格'
                elif time.monotonic() > deadline: reason = f'读取时间超过 {seconds} 秒'
                if reason:
                    recovery_warnings.append(f'部分入库：{reason}；从 {sheet.title} 第 {row_no} 行起未读取，可开启大文件模式后重新同步。')
                    stopped = True; break
                count += len(row)
                if not row: continue
                if progress and row_no % 200 == 0: progress(stage='reading', location=f'{sheet.title} · 第 {row_no} 行')
                if not context:
                    context = '首个非空行参考：' + ' | '.join(f'{get_column_letter(i)}={formula_text(c.value) if c.data_type == "f" else c.value}' for i,c in row.items())[:2000]
                items, cells = [], {}
                for col, cell in row.items():
                    cached_cell = values.get(col)
                    display = cached_cell.value if cell.data_type == 'f' and cached_cell else cell.value
                    if cell.data_type == 'f' and cached_cell is None: display = '未保存（未重新计算）'
                    if isinstance(display, datetime):
                        display = display.date().isoformat() if display.time() == dt_time() and 'h' not in cell.number_format.lower() and 's' not in cell.number_format.lower() else display.isoformat(sep=' ')
                    elif isinstance(display, (date, dt_time)): display = display.isoformat()
                    extras = {'coordinate': cell.coordinate, 'number_format': cell.number_format}
                    if cell.data_type == 'f':
                        extras['formula'] = formula_text(cell.value)
                        extras['cached_result'] = cached_cell is not None
                    cells[get_column_letter(col)] = table_cell(display, **extras)
                    items.append(f'{cell.coordinate}={display}')
                line = f'[{sheet.title} 行 {row_no}] ' + ' | '.join(items)
                if parts and size + len(line) + 1 > UNIT_SIZE: flush()
                if text_count + len(line) + len(context) + 1 > text_limit or len(units) >= unit_limit:
                    recovery_warnings.append(f'部分入库：达到文本或阅读单元上限；从 {sheet.title} 第 {row_no} 行起未读取，可开启大文件模式后重新同步。')
                    stopped = True; break
                if not parts:
                    parts = [context]; size = len(context)+1; text_count += size
                first = first or row_no; last = row_no
                grid_rows.append({'index': row_no, 'cells': cells, '_start': size, '_end': size+len(line)})
                parts.append(line); size += len(line)+1; text_count += len(line)+1
            flush()
            if stopped: break
        return units, props, recovery_warnings + ['公式不执行；检索采用保存的结果，原公式保留在表格详情中。保存结果可能缺失或过期。']
    except BaseException:
        units.close(); raise
    finally:
        for book in (formula_book, cached_book):
            if book:
                book.close()
                if hasattr(book, '_research_stream'): book._research_stream.close()


def delimited_units(text, delimiter, options=None, cancel=None):
    units, parts, size, first, last, header = Spool(), [], 0, None, None, None
    grid_rows = []
    reader = csv.reader(io.StringIO(text, newline=''), delimiter=delimiter)
    total = 0
    def flush():
        nonlocal parts, size, first, last, grid_rows
        if parts:
            units.append(Unit('\n'.join(parts), {'kind': 'records', 'record_start': first, 'record_end': last,
                'label': f'记录 {first}–{last}'}, tables=[make_table('表格', grid_rows,
                    [{'key': str(i+1), 'label': f'{i+1} · {header[i]}' if i < len(header) and header[i] else str(i+1)}
                     for i in range(max(len(r['cells']) for r in grid_rows))], '列标题参考第一条记录')]))
        parts, size, first, last, grid_rows = [], 0, None, None, []
    for n, row in enumerate(reader, 1):
        from office_ocr import cancelled
        cancelled(cancel)
        total += sum(bool(v) for v in row)
        if total > reading_limits(options)[0]: raise ValueError('表格单元格数超过读取限制。')
        if header is None: header = row
        line = f'[记录 {n}] ' + ' | '.join(f'列 {i+1}（{header[i] if i < len(header) else ""}）={v!r}' for i, v in enumerate(row))
        if size + len(line) > UNIT_SIZE: flush()
        first = first or n; last = n
        grid_rows.append({'index': n, 'cells': {str(i+1): table_cell(v) for i, v in enumerate(row)}, '_start': size, '_end': size + len(line)})
        parts.append(line); size += len(line) + 1
    flush()
    return units

def extract_document(path, progress=None, cancel=None, options=None):
    original_suffix = suffix = path.suffix.lower()
    format_warning = None
    if suffix == '.xls' and zipfile.is_zipfile(path):
        check_package(path)
        with zipfile.ZipFile(path) as z:
            if 'xl/workbook.xml' not in z.namelist():
                raise ValueError('此 XLS 扩展名对应的内容不是 XLSX 工作簿；请检查实际文件格式。')
        suffix = '.xlsx'
        format_warning = '已按文件内容识别为 XLSX；原文件使用 XLS 扩展名，文件名及内容未修改。'

    if suffix == '.xls':
        from text_workbook import detect, converted
        kind = detect(path)
        if kind:
            with converted(path, kind) as derived:
                metadata, units = extract_document(derived, progress, cancel, options)
            metadata.update(file_name=path.name, file_format='xls', detected_format=kind)
            if metadata['metadata_provenance']['fields']['title']['source'] == 'filename':
                metadata['title'] = path.stem
                metadata['metadata_provenance']['fields']['title']['evidence'] = path.stem
            metadata['extraction_warnings'].append('按实际内容识别为文本表格；只读取表格文字，不运行脚本、宏或外部链接。位置按提取表格计算。')
            return metadata, units
    if suffix in LEGACY:
        from legacy_office import convert
        derived, conversion = convert(path, progress, cancel)
        metadata, units = extract_document(derived, progress, cancel, options)
        metadata.update(file_name=path.name, file_format=suffix[1:], conversion=conversion)
        fields = metadata['metadata_provenance']['fields']
        if fields['title']['source'] == 'filename':
            metadata['title'] = path.stem
            fields['title']['evidence'] = path.stem
        warning = '旧版 Office 经本地转换后提取；位置以转换副本为准，排版或公式缓存可能变化，请核对原文件。'
        metadata['extraction_warnings'].append(warning)
        converted_units = Spool()
        for unit in units:
            unit.location.update(source_format=suffix[1:], location_basis='converted_copy')
            unit.location['label'] += '（转换副本）'
            unit.warnings.append(warning)
            converted_units.append(unit)
        return metadata, converted_units
    if suffix not in SUPPORTED - {'.pdf'}: raise ValueError('不支持此文件格式')
    if path.stat().st_size > MAX_BYTES: raise ValueError('非 PDF 文件超过 100 MB 读取限制。')
    props, warnings, encoding, ocr_summary = {}, [], None, None
    if suffix in {'.docx', '.pptx', '.xlsx'}:
        check_package(path)
        units, props, warnings = spreadsheet_units(path, options, cancel, progress) if suffix == '.xlsx' else {'.docx': word_units, '.pptx': presentation_units}[suffix](path)
        from office_ocr import extract_images
        images, failures, ocr_summary = extract_images(path, progress, cancel, suffix)
        units.extend(images)
        warnings.extend(failures)
        if failures: warnings.append('部分入库：部分内嵌图片未完成 OCR；已保留正文及其他可识别内容。')
        if ocr_summary['images']:
            warnings.append('内嵌图片使用本地 OCR；文字和数字可能有误，不还原图表关系。图片预览保留出处，但不代表整页排版。')
        if not any(u.text.strip() for u in units) and failures:
            raise ValueError('未提取到可检索文字；' + '；'.join(failures[:3]))
    else:
        raw = path.read_bytes()
        if suffix == '.rtf':
            from striprtf.striprtf import rtf_to_text
            text = rtf_to_text(raw.decode('latin-1'))
            units = text_units(text, 'text_lines')
            warnings = ['RTF 行号指提取后的正文行；图片和排版未保留。']
        else:
            text, encoding = decode_text(raw)
            units = delimited_units(text, '\t' if suffix == '.tsv' else ',', options, cancel) if suffix in {'.csv', '.tsv'} else text_units(text)
    if format_warning: warnings.append(format_warning)
    _, text_limit, unit_limit, _ = reading_limits(options)
    if not any(u.text.strip() for u in units):
        raise ValueError('未提取到可检索文字；文档为空、图片没有可识别文字，或图片格式不受支持。')
    # Bound response size even for unusually long source lines/rows.
    bounded = Spool()
    text_count = 0
    for u in units:
        if text_count + len(u.text) > text_limit or len(bounded) >= unit_limit:
            warnings.append(f'部分入库：达到 {text_limit:,} 字符或 {unit_limit:,} 个阅读单元上限；其余内容未入库，可开启大文件模式后重新同步。')
            break
        text_count += len(u.text)
        start = 0
        while start < max(1, len(u.text)):
            end = min(start + UNIT_SIZE, len(u.text))
            if 0 < len(u.text) - end < 256:
                end = len(u.text) - 256
            loc = dict(u.location)
            if len(u.text) > UNIT_SIZE:
                part_number = 1 if start == 0 else (start + 255)//UNIT_SIZE+1
                loc.update(part=part_number, character_start=start+1, character_end=end)
                loc['label'] += f' · 片段 {part_number}'
            bounded.append(Unit(u.text[start:end], loc, u.warnings, slice_tables(u.tables, start, end)))
            if end == len(u.text): break
            start = end
    if isinstance(units, Spool): units.close()
    units = bounded
    if not units: raise ValueError('首个阅读单元超过当前限额；请开启大文件模式或拆分文件。')
    title = (props.get('title') or '').strip() or path.stem
    cover = units[0].text
    # Only unprefixed explicit lines count as publisher/date evidence.
    evidence = re.sub(r'^\[段落 \d+\] ', '', cover, flags=re.M)
    institution, inst_evidence = institution_from_cover(evidence)
    published, display, precision, date_evidence = cover_date(evidence)
    fields = {'title': {'source': 'embedded_metadata' if props.get('title') else 'filename', 'evidence': title}}
    if institution: fields['institution'] = {'source': 'first_unit', 'evidence': inst_evidence, 'needs_review': units[0].location.get('kind') == 'embedded_image'}
    if display: fields['published_date'] = {'source': 'first_unit', 'evidence': date_evidence, 'needs_review': units[0].location.get('kind') == 'embedded_image'}
    metadata = {'title': title, 'file_name': path.name, 'file_format': suffix[1:], 'pdf_title': None,
        'author': props.get('author') or None, 'institution': institution, 'published_date': published,
        'publication_date_display': display, 'publication_date_precision': precision,
        'source_type': 'unclassified', 'companies': [], 'metadata_schema_version': SCHEMA_VERSION,
        'parser_version': PARSER_VERSION, 'large_file_mode': (options or {}).get('large_file_mode', False),
        'partial': any('部分入库' in w for w in warnings), 'detected_format': suffix[1:], 'ocr_summary': ocr_summary, 'encoding': encoding, 'extraction_warnings': warnings,
        'metadata_provenance': {'fields': fields, 'policy': 'Explicit source evidence only; creation dates are not publication dates.'}}
    return metadata, units
